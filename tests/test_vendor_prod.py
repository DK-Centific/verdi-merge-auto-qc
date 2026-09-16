"""Rate Approval Log mapping, dedupe, and local append."""
from __future__ import annotations

import io
import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import ots_pipeline as pipe


def _sample_master_bytes(*, existing_name: str = "already_there.wav") -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = pipe.RATE_APPROVAL_SHEET
    headers = [""] * 22
    headers[0] = "Vendor Name"
    headers[1] = "Project Code"
    headers[4] = "Workflow"
    headers[5] = "Locale"
    headers[6] = "Vendor's file submit date (Ingestion Batch)"
    headers[12] = "Qualification Rate (Pass/Submitted)"
    headers[14] = "Timeline payout"
    headers[19] = "File Name"
    headers[20] = "File path"
    headers[21] = "Durations (seconds)"
    ws.append(headers)
    ws.append(
        [
            "alchemy",
            "Maple",
            "",
            "",
            "Maple",
            "el-GR",
            datetime(2026, 8, 23),
            "",
            "",
            "",
            "",
            "",
            '=IFERROR(L2/K2, "")',
            "",
            "",
            "",
            "",
            "",
            "",
            existing_name,
            "alchemyhandoff/maple/el-GR/" + existing_name,
            12.5,
        ]
    )
    ws.append(["", "", "", "", "", "", "", "", "", "", "", "", "", "", "Notes — keep below", "", "", "", "", "", "", ""])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class RateApprovalMappingTests(unittest.TestCase):
    def test_to_rate_approval_rows_uses_enriched_fields(self):
        rows = pipe.to_rate_approval_rows(
            [
                {
                    "vendorName": "Alchemy",
                    "_vendor": "ignored",
                    "projectCode": "Maple",
                    "workflow": "Maple",
                    "locale": "el-GR",
                    "ingestionBatch": "2026-08-23",
                    "fileName": "clip.wav",
                    "filePath": "alchemy/clip.wav",
                    "durationSeconds": 9.5,
                }
            ]
        )
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["vendorName"], "Alchemy")
        self.assertEqual(rows[0]["fileName"], "clip.wav")
        self.assertEqual(rows[0]["durationSeconds"], 9.5)
        self.assertEqual(rows[0]["locale"], "el-GR")

    def test_classify_file_name_is_primary_key(self):
        existing_names = {"clip.wav"}
        classified = pipe.classify_rate_approval_rows(
            [
                {
                    "vendorName": "Alchemy",
                    "projectCode": "Maple",
                    "workflow": "Maple",
                    "locale": "el-GR",
                    "fileName": "CLIP.WAV",
                    "filePath": "other/path.wav",
                    "durationSeconds": 1,
                },
                {
                    "vendorName": "Alchemy",
                    "projectCode": "Maple",
                    "workflow": "Maple",
                    "locale": "el-GR",
                    "fileName": "new.wav",
                    "filePath": "p/new.wav",
                    "durationSeconds": 2,
                },
                {
                    "vendorName": "",
                    "fileName": "no-vendor.wav",
                },
            ],
            existing_names,
            set(),
        )
        self.assertEqual(classified["duplicate_count"], 1)
        self.assertEqual(classified["unique_count"], 1)
        self.assertEqual(classified["incomplete_count"], 1)
        self.assertEqual(classified["unique_rows"][0]["fileName"], "new.wav")

    def test_intra_batch_duplicate_file_name(self):
        classified = pipe.classify_rate_approval_rows(
            [
                {"vendorName": "A", "projectCode": "Maple", "locale": "el-GR", "fileName": "same.wav", "filePath": "a"},
                {"vendorName": "B", "projectCode": "Maple", "locale": "he-IL", "fileName": "same.wav", "filePath": "b"},
            ],
            set(),
            set(),
        )
        self.assertEqual(classified["unique_count"], 1)
        self.assertEqual(classified["duplicate_count"], 1)


class AppendRateApprovalTests(unittest.TestCase):
    def test_append_bytes_only_unique_and_preserves_formula(self):
        original = _sample_master_bytes()
        new_bytes, summary = pipe.append_rate_approval_log_bytes(
            original,
            [
                {
                    "vendorName": "alchemy",
                    "projectCode": "Maple",
                    "workflow": "Maple",
                    "locale": "el-GR",
                    "ingestionBatch": "2026-08-23",
                    "fileName": "already_there.wav",
                    "filePath": "x",
                    "durationSeconds": 99,
                },
                {
                    "vendorName": "pangeanic",
                    "projectCode": "Maple",
                    "workflow": "Maple",
                    "locale": "es-ES",
                    "ingestionBatch": "2026-09-01",
                    "fileName": "brand_new.wav",
                    "filePath": "pangeanic/brand_new.wav",
                    "durationSeconds": 3.25,
                },
            ],
        )
        self.assertEqual(summary["appended"], 1)
        self.assertEqual(summary["skipped_duplicates"], 1)
        wb = openpyxl.load_workbook(io.BytesIO(new_bytes))
        ws = wb[pipe.RATE_APPROVAL_SHEET]
        # Existing formula in M2 must remain
        self.assertEqual(ws.cell(2, 13).value, '=IFERROR(L2/K2, "")')
        # New unique row written after existing data, notes shifted down
        names = [str(ws.cell(r, 20).value or "") for r in range(2, ws.max_row + 1)]
        self.assertIn("already_there.wav", names)
        self.assertIn("brand_new.wav", names)
        new_row = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, 20).value == "brand_new.wav")
        self.assertEqual(ws.cell(new_row, 1).value, "pangeanic")
        self.assertEqual(ws.cell(new_row, 2).value, "Maple")
        self.assertEqual(ws.cell(new_row, 5).value, "Maple")
        self.assertEqual(ws.cell(new_row, 6).value, "es-ES")
        self.assertEqual(ws.cell(new_row, 22).value, 3.25)
        # Do not invent values in unused columns on the new row
        self.assertIn(ws.cell(new_row, 3).value, (None, ""))
        notes = [str(ws.cell(r, 15).value or "") for r in range(1, ws.max_row + 1)]
        self.assertTrue(any("Notes" in n for n in notes))

    def test_append_path_roundtrip(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "Vendor Prod Files.xlsx"
            path.write_bytes(_sample_master_bytes())
            summary = pipe.append_rate_approval_log(
                path,
                [
                    {
                        "vendorName": "alchemy",
                        "projectCode": "Maple",
                        "locale": "el-GR",
                        "fileName": "second.wav",
                        "filePath": "p/second.wav",
                        "durationSeconds": 1,
                    }
                ],
            )
            self.assertEqual(summary["appended"], 1)
            names, _ = pipe.extract_existing_rate_keys_from_path(path)
            self.assertIn("already_there.wav", names)
            self.assertIn("second.wav", names)


class PreviewAndWorkbookTests(unittest.TestCase):
    def test_preview_and_result_sheet(self):
        preview = pipe.preview_vendor_prod(
            [
                {
                    "vendorName": "Alchemy",
                    "projectCode": "Maple",
                    "workflow": "Maple",
                    "locale": "el-GR",
                    "ingestionBatch": "2026-08-23",
                    "fileName": "clip.wav",
                    "filePath": "a/clip.wav",
                    "durationSeconds": 4,
                    "_vendor": "Alchemy",
                }
            ]
        )
        self.assertEqual(preview["unique_count"], 1)
        xlsx = pipe.write_vendor_prod_preview_xlsx(preview)
        wb = openpyxl.load_workbook(io.BytesIO(xlsx))
        self.assertIn("Vendor_Prod_Rows", wb.sheetnames)
        self.assertEqual(wb["Vendor_Prod_Rows"].cell(2, 6).value, "clip.wav")

    def test_write_result_includes_vendor_prod_sheet(self):
        annotated = [
            {
                "status": "",
                "feedback": "",
                "folderName": "f",
                "fileName": "clip.wav",
                "locale": "el-GR",
                "vendorName": "Alchemy",
                "projectCode": "Maple",
                "workflow": "Maple",
                "ingestionBatch": "2026-08-23",
                "filePath": "f/clip.wav",
                "durationSeconds": 4,
                "_vendor": "Alchemy",
            }
        ]
        preview = pipe.preview_vendor_prod(annotated)
        xlsx = pipe.write_result_workbook(annotated, [], [], [], vendor_prod_preview=preview)
        wb = openpyxl.load_workbook(io.BytesIO(xlsx))
        self.assertIn("Vendor_Prod_Rows", wb.sheetnames)
        self.assertIn("Merged", wb.sheetnames)


class UploadedSchemaTests(unittest.TestCase):
    """Optional: local upload of a real Vendor Prod copy (never committed)."""

    def test_extract_keys_from_uploaded_prod_copy_if_present(self):
        uploaded = Path("/home/ubuntu/.cursor/projects/workspace/uploads/Vendor_Prod_Files_cde7.xlsx")
        if not uploaded.exists():
            self.skipTest("uploaded prod copy not present")
        names, composites = pipe.extract_existing_rate_keys_from_path(uploaded)
        self.assertGreater(len(names), 10)
        wb = openpyxl.load_workbook(uploaded, read_only=True)
        self.assertIn(pipe.RATE_APPROVAL_SHEET, wb.sheetnames)
        ws = wb[pipe.RATE_APPROVAL_SHEET]
        self.assertEqual(str(ws.cell(1, 1).value), "Vendor Name")
        self.assertEqual(str(ws.cell(1, 20).value), "File Name")
        self.assertEqual(str(ws.cell(1, 22).value), "Durations (seconds)")


if __name__ == "__main__":
    unittest.main()
